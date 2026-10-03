"""仅在内存中解析获准资料，不执行或联网。 / Parse approved bytes without execution or network I/O."""

from __future__ import annotations

import csv
import io
import json
import ntpath
import re
import stat
import threading
import zipfile
from pathlib import PurePosixPath
from typing import Any, Iterable


MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_TEXT_CHARACTERS = 150_000
MAX_TABLE_ROWS = 3_000
MAX_TABLE_COLUMNS = 200
MAX_ZIP_MEMBERS = 100
MAX_ZIP_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_ZIP_COMPRESSION_RATIO = 100
MAX_PDF_PAGES = 200
TRUNCATION_MARKER = "[提取内容已截断；原件仍完整保留，请按需拆分材料。]"
_CSV_LIMIT_LOCK = threading.RLock()


def _result(status: str, error: str = "") -> dict:
    return {"status": status, "text": "", "segments": [], "tables": [], "error": error}


def _safe_member_name(info: zipfile.ZipInfo) -> str:
    # 同时检查 Windows 和 POSIX 路径，防止后续归档层受到路径穿越影响。
    # Validate both Windows and POSIX paths before any archive consumer sees a name.
    original = getattr(info, "orig_filename", info.filename)
    filename = original.replace("\\", "/")
    if "\x00" in original or not filename or filename.startswith("/"):
        raise ValueError("资料包包含空名称、绝对路径或非法字符。")
    if ntpath.splitdrive(filename)[0] or any(":" in part for part in filename.split("/")):
        raise ValueError("资料包包含磁盘路径或备用数据流路径。")
    parts = filename.rstrip("/").split("/")
    if any(part in ("", ".", "..") or part.rstrip(" .") != part for part in parts):
        raise ValueError("资料包包含目录穿越、空路径或不安全的文件名。")
    if any(part.upper().split(".", 1)[0] in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))} for part in parts):
        raise ValueError("资料包包含系统保留文件名。")
    mode = info.external_attr >> 16
    if stat.S_ISLNK(mode) or (stat.S_IFMT(mode) and not (stat.S_ISREG(mode) or stat.S_ISDIR(mode))):
        raise ValueError("资料包包含符号链接或特殊文件。")
    if info.flag_bits & 1:
        raise ValueError("资料包已加密，未尝试解密。")
    return str(PurePosixPath(filename))


def safe_zip_members(raw: bytes) -> list[dict]:
    """校验并读取一层 ZIP，任何安全限制失败均拒绝整包。 / Validate and read one ZIP level atomically."""
    if not isinstance(raw, bytes):
        raise ValueError("资料包必须是字节内容。")
    if len(raw) > MAX_FILE_BYTES:
        raise ValueError("资料包超过单文件容量限制（25 MB）。")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_MEMBERS:
                raise ValueError(f"资料包条目超过 {MAX_ZIP_MEMBERS} 项，未展开。")
            metadata = []
            total_size = 0
            seen = set()
            for info in entries:
                filename = _safe_member_name(info)
                identity = filename.casefold()
                if identity in seen:
                    raise ValueError("资料包包含重复路径，请先区分名称后重新导入。")
                seen.add(identity)
                if info.file_size > MAX_FILE_BYTES:
                    raise ValueError("资料包包含超过单项容量限制（25 MB）的文件。")
                if info.file_size / max(info.compress_size, 1) > MAX_ZIP_COMPRESSION_RATIO:
                    raise ValueError(f"资料包压缩比超过安全限制（{MAX_ZIP_COMPRESSION_RATIO} 倍）。")
                total_size += info.file_size
                if total_size > MAX_ZIP_EXPANDED_BYTES:
                    raise ValueError("资料包展开总量超过安全限制（100 MB）。")
                if not info.is_dir():
                    metadata.append((info, filename))

            members = []
            actual_total = 0
            for info, filename in metadata:
                # 分块读取并再次限制实际大小，不只相信 ZIP 元数据。
                # Bound actual decompressed bytes as well as ZIP metadata.
                chunks = []
                actual_size = 0
                with archive.open(info) as source:
                    while chunk := source.read(64 * 1024):
                        actual_size += len(chunk)
                        actual_total += len(chunk)
                        if actual_size > MAX_FILE_BYTES or actual_total > MAX_ZIP_EXPANDED_BYTES:
                            raise ValueError("资料包实际展开大小超过安全限制。")
                        chunks.append(chunk)
                if actual_size != info.file_size:
                    raise ValueError("资料包文件长度与元数据不一致。")
                members.append({"filename": filename, "content": b"".join(chunks), "error": ""})
            return members
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("资料包损坏或使用了不支持的压缩方式，未展开。") from exc


def _check_office_archive(raw: bytes) -> list[dict]:
    members = safe_zip_members(raw)
    for member in members:
        if member["filename"].lower().endswith((".xml", ".rels")):
            # 去除零字节也能识别 UTF-16/32 声明，不允许 DTD 或实体扩展。
            # Removing NULs detects UTF-16/32 declarations; DTD/entity expansion is disallowed.
            content = member["content"].replace(b"\x00", b"")
            if re.search(rb"<!\s*(DOCTYPE|ENTITY)\b", content, flags=re.IGNORECASE):
                raise ValueError("Office 文件含 DTD 或实体声明，出于安全原因不解析。")
    return members


def _decode_text(raw: bytes) -> str:
    if b"\x00" in raw:
        raise ValueError("文本包含二进制零字节；请另存为 UTF-8 或 GB18030 文本。")
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文本编码；请另存为 UTF-8 或 GB18030。")


class _Collector:
    def __init__(self) -> None:
        self.segments: list[dict] = []
        self.tables: list[dict] = []
        self.notes: list[str] = []
        self.characters = 0
        self.table_rows = 0
        self.truncated = False

    @property
    def remaining(self) -> int:
        return max(0, MAX_TEXT_CHARACTERS - len(TRUNCATION_MARKER) - 2 - self.characters)

    def note(self, message: str) -> None:
        if message not in self.notes:
            self.notes.append(message)

    def truncate(self, message: str) -> None:
        self.truncated = True
        self.note(message)

    def add(self, anchor: str, content: str, *, strip: bool = True) -> str:
        content = content.strip() if strip else content
        if not content.strip():
            return ""
        separator = 2 if self.segments else 0
        available = max(0, self.remaining - separator)
        retained = content[:available]
        if len(retained) < len(content):
            self.truncate(f"正文提取达到 {MAX_TEXT_CHARACTERS} 字符上限。")
        if retained:
            self.segments.append({"anchor": anchor, "text": retained})
            self.characters += len(retained) + separator
        return retained

    def finish(self) -> dict:
        content = "\n\n".join(segment["text"] for segment in self.segments)
        if self.truncated:
            content = f"{content}\n\n{TRUNCATION_MARKER}".lstrip()
        return {"status": "parsed", "text": content, "segments": self.segments, "tables": self.tables, "error": " ".join(self.notes)}


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value).replace("\t", " ")


def _row(collector: _Collector, anchor: str, values: list[str]) -> list[str]:
    # 表格与文本使用同一截断结果，不让结构化数据绕开提取上限。
    # Tables and excerpts share the same clipped values and extraction budget.
    content = "\t".join(values)
    if not content.strip():
        return [""] * len(values)
    retained = collector.add(anchor, content, strip=False)
    clipped = retained.split("\t") if retained else []
    return clipped + [""] * (len(values) - len(clipped))


def _add_table(collector: _Collector, name: str, raw_rows: Iterable[tuple[int, Iterable[Any]]]) -> None:
    iterator = iter(raw_rows)
    header = None
    for position, values in iterator:
        cells = [_cell(value) for value in values]
        while cells and not cells[-1]:
            cells.pop()
        if cells and any(cell.strip() for cell in cells):
            header = cells
            break
    if header is None:
        return
    collector.add(f"{name} · 表名", name)
    if len(header) > MAX_TABLE_COLUMNS:
        collector.truncate(f"表格仅提取前 {MAX_TABLE_COLUMNS} 列。")
    columns = _row(collector, f"{name} · 第 {position} 行（表头）", header[:MAX_TABLE_COLUMNS])
    columns = [value or f"列{index + 1}" for index, value in enumerate(columns)]
    rows = []
    for position, values in iterator:
        cells = [_cell(value) for value in values]
        while cells and not cells[-1]:
            cells.pop()
        if not cells or not any(cell.strip() for cell in cells):
            continue
        if collector.table_rows >= MAX_TABLE_ROWS:
            collector.truncate(f"表格数据合计仅提取前 {MAX_TABLE_ROWS} 行。")
            break
        if not collector.remaining:
            collector.truncate(f"正文提取达到 {MAX_TEXT_CHARACTERS} 字符上限。")
            break
        if len(cells) > MAX_TABLE_COLUMNS:
            collector.truncate(f"表格仅提取前 {MAX_TABLE_COLUMNS} 列。")
        cells = cells[:MAX_TABLE_COLUMNS]
        if len(cells) > len(columns):
            previous_width = len(columns)
            columns.extend(f"列{index + 1}" for index in range(previous_width, len(cells)))
            for previous in rows:
                previous.extend([""] * (len(columns) - len(previous)))
        cells.extend([""] * (len(columns) - len(cells)))
        rows.append(_row(collector, f"{name} · 第 {position} 行", cells))
        collector.table_rows += 1
    collector.tables.append({"name": name, "columns": columns, "rows": rows})


def _plain_text(raw: bytes, collector: _Collector) -> None:
    for index, line in enumerate(_decode_text(raw).splitlines(), start=1):
        if not collector.remaining and line.strip():
            collector.truncate(f"正文提取达到 {MAX_TEXT_CHARACTERS} 字符上限。")
            break
        collector.add(f"第 {index} 行", line)
        if not collector.remaining and collector.truncated:
            break


def _csv(raw: bytes, collector: _Collector) -> None:
    content = _decode_text(raw)
    try:
        dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    # CSV 默认字段限制过小；暂时提高并在解析后恢复进程设置。
    # Temporarily raise CSV's field limit and restore the process setting afterwards.
    with _CSV_LIMIT_LOCK:
        previous_limit = csv.field_size_limit()
        try:
            csv.field_size_limit(MAX_FILE_BYTES)
            rows = csv.reader(io.StringIO(content, newline=""), dialect=dialect, strict=True)
            _add_table(collector, "CSV", enumerate(rows, start=1))
        finally:
            csv.field_size_limit(previous_limit)


def _json_table(value: Any, collector: _Collector, name: str) -> bool:
    if not isinstance(value, list) or not value or not all(isinstance(item, dict) for item in value[:MAX_TABLE_ROWS + 1]):
        return False
    columns = list(dict.fromkeys(key for item in value[:MAX_TABLE_ROWS + 1] for key in item))
    if not columns:
        return False
    if len(columns) > MAX_TABLE_COLUMNS:
        collector.truncate(f"表格仅提取前 {MAX_TABLE_COLUMNS} 列。")
        columns = columns[:MAX_TABLE_COLUMNS]

    def rows() -> Iterable[tuple[int, Iterable[Any]]]:
        yield 1, columns
        for index, item in enumerate(value, start=2):
            if not isinstance(item, dict):
                collector.note("JSON 数组中后续结构不一致的条目未作为表格行解释。")
                continue
            yield index, [item.get(column, "") for column in columns]

    _add_table(collector, name, rows())
    return True


def _json(raw: bytes, collector: _Collector) -> None:
    value = json.loads(_decode_text(raw))
    if _json_table(value, collector, "JSON"):
        return
    if isinstance(value, dict):
        for key, item in value.items():
            if not collector.remaining:
                collector.truncate(f"正文提取达到 {MAX_TEXT_CHARACTERS} 字符上限。")
                break
            if not _json_table(item, collector, f"JSON /{key}"):
                collector.add(f"JSON /{key}", f"{key}: {json.dumps(item, ensure_ascii=False, indent=2)}")
    else:
        collector.add("JSON 根节点", json.dumps(value, ensure_ascii=False, indent=2))


def _docx(raw: bytes, collector: _Collector) -> None:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_office_archive(raw)
    document = Document(io.BytesIO(raw))
    paragraph_index = 0
    table_index = 0
    for element in document.element.body:
        if element.tag == qn("w:p"):
            paragraph_index += 1
            collector.add(f"段落 {paragraph_index}", Paragraph(element, document).text)
        elif element.tag == qn("w:tbl"):
            table_index += 1
            table = Table(element, document)
            _add_table(collector, f"表格 {table_index}", ((index, (cell.text for cell in row.cells)) for index, row in enumerate(table.rows, start=1)))
        if not collector.remaining and collector.truncated:
            break
    collector.note("已提取正文段落和表格；图片、批注、页眉页脚及嵌入对象未分析。")


def _xlsx(raw: bytes, collector: _Collector) -> None:
    from openpyxl import load_workbook

    _check_office_archive(raw)
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    try:
        for sheet in workbook.worksheets:
            if not collector.remaining:
                collector.truncate(f"正文提取达到 {MAX_TEXT_CHARACTERS} 字符上限。")
                break
            if sheet.max_column and sheet.max_column > MAX_TABLE_COLUMNS:
                collector.truncate(f"表格仅提取前 {MAX_TABLE_COLUMNS} 列。")
            if sheet.max_row and sheet.max_row > MAX_TABLE_ROWS + 2:
                collector.truncate(f"工作表只检查前 {MAX_TABLE_ROWS + 2} 行，并最多提取 {MAX_TABLE_ROWS} 行数据。")
            rows = sheet.iter_rows(values_only=True, max_row=MAX_TABLE_ROWS + 2, max_col=MAX_TABLE_COLUMNS)
            _add_table(collector, f"工作表 {sheet.title}", enumerate(rows, start=1))
    finally:
        workbook.close()
    collector.note("公式未执行，仅读取文件已有缓存值；无缓存结果留空，图片和图表未分析。")


def _pdf(raw: bytes, collector: _Collector) -> None:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw), strict=True)
    if reader.is_encrypted:
        raise ValueError("PDF 已加密，未尝试解密。")
    page_count = len(reader.pages)
    if page_count > MAX_PDF_PAGES:
        collector.truncate(f"PDF 仅提取前 {MAX_PDF_PAGES} 页。")
    empty_pages = 0
    for index in range(min(page_count, MAX_PDF_PAGES)):
        if not collector.remaining:
            collector.truncate(f"正文提取达到 {MAX_TEXT_CHARACTERS} 字符上限。")
            break
        content = reader.pages[index].extract_text() or ""
        if not content.strip():
            empty_pages += 1
        collector.add(f"第 {index + 1} 页", content)
    if not collector.segments:
        raise ValueError("PDF 未提取到文字，可能是扫描件或空白页；原件保留，当前不支持 OCR。")
    if empty_pages:
        collector.note(f"有 {empty_pages} 页未提取到文字；图片和扫描内容未进行 OCR。")
    collector.note("PDF 按页提取文字；复杂布局、图片与表格结构未解析。")


def parse_material(filename: str, raw: bytes) -> dict:
    """返回归档资料的解析状态与可定位内容。 / Return parse status and source-located content."""
    if not isinstance(raw, bytes):
        return _result("failed", "解析输入必须是字节内容。")
    if len(raw) > MAX_FILE_BYTES:
        return _result("failed", "文件超过单文件容量限制（25 MB），未解析。")
    extension = PurePosixPath(str(filename).replace("\\", "/")).suffix.lower()
    handlers = {".txt": _plain_text, ".md": _plain_text, ".csv": _csv, ".json": _json, ".docx": _docx, ".xlsx": _xlsx, ".pdf": _pdf}
    if extension not in handlers:
        if extension in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff", ".svg"}:
            explanation = "图片仅归档，未进行 OCR 或图片理解。"
        elif extension in {".mp3", ".wav", ".m4a", ".flac", ".ogg", ".mp4", ".mov", ".avi", ".mkv", ".webm"}:
            explanation = "音视频仅归档，未转写、观看或分析内容。"
        elif extension == ".zip":
            explanation = "ZIP 原件仅归档；包内文件需要经过安全检查后逐项解析，不递归展开。"
        else:
            explanation = "当前格式仅归档，尚未支持内容解析；文件不会被执行。"
        return _result("unsupported", explanation)
    collector = _Collector()
    try:
        handlers[extension](raw, collector)
        result = collector.finish()
        if not result["text"]:
            return _result("failed", "未提取到可用文字或表格；原件仍保留。")
        return result
    except ImportError:
        return _result("failed", "本地解析依赖尚未安装；原件保留，可安装后重试。")
    except (ValueError, csv.Error) as exc:
        return _result("failed", str(exc))
    except Exception:
        # 解析器报错不暴露堆栈、服务器路径或文档私密内容。
        # Keep library exceptions from exposing stack traces, server paths, or document contents.
        return _result("failed", "文件损坏、结构不受支持或无法安全解析；原件保留，可重试或另存支持格式。")
