"""用内存中的虚构资料验证解析与归档安全。 / Test parser safety with synthetic in-memory materials."""

import io
import json
import stat
import unittest
import zipfile
from unittest.mock import patch

from operations.parsers import MAX_TEXT_CHARACTERS, parse_material, safe_zip_members


def _zip(entries, *, compression=zipfile.ZIP_STORED):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=compression) as archive:
        for name, content in entries:
            archive.writestr(name, content)
    return output.getvalue()


def _pdf(text="Synthetic workshop evidence"):
    # 创建最小合法 PDF，无需磁盘、外部字体或第三方生成器。
    # Build a minimal valid PDF without disk access, external fonts, or a generator dependency.
    stream = f"BT /F1 12 Tf 20 180 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, content in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii") + content + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii"))
    return bytes(output)


class TextParserTests(unittest.TestCase):
    def test_utf8_text_keeps_line_anchors(self):
        result = parse_material("活动记录.TXT", "活动完成\n\n下一步修正设备检查".encode())
        self.assertEqual(result["status"], "parsed")
        self.assertEqual(result["segments"][1]["anchor"], "第 3 行")
        self.assertIn("设备检查", result["text"])

    def test_utf8_bom_markdown_and_gb18030(self):
        for encoding in ("utf-8-sig", "gb18030"):
            with self.subTest(encoding=encoding):
                result = parse_material("说明.md", "# 虚构活动\n作品说明".encode(encoding))
                self.assertEqual(result["status"], "parsed")
                self.assertIn("虚构活动", result["text"])
                self.assertNotIn("\ufeff", result["text"])

    def test_empty_and_binary_text_are_not_success(self):
        for raw in (b"", b" \n\r", b"a\x00b", b"\xff"):
            with self.subTest(raw=raw):
                self.assertEqual(parse_material("empty.txt", raw)["status"], "failed")

    def test_text_limit_is_explicit_and_bounded(self):
        result = parse_material("large.txt", ("虚构" * MAX_TEXT_CHARACTERS).encode())
        self.assertEqual(result["status"], "parsed")
        self.assertLessEqual(len(result["text"]), MAX_TEXT_CHARACTERS)
        self.assertLessEqual(sum(len(segment["text"]) for segment in result["segments"]), MAX_TEXT_CHARACTERS)
        self.assertIn("截断", result["text"])
        self.assertIn("字符上限", result["error"])

    def test_input_size_is_checked_before_parse(self):
        with patch("operations.parsers.MAX_FILE_BYTES", 4):
            self.assertEqual(parse_material("large.txt", b"12345")["status"], "failed")

    def test_unsupported_files_are_archive_only(self):
        for name in ("photo.png", "clip.mp4", "sound.wav", "slides.pptx", "script.exe", "old.doc", "archive.zip"):
            with self.subTest(name=name):
                result = parse_material(name, b"not-executed")
                self.assertEqual(result["status"], "unsupported")
                self.assertIn("归档", result["error"])
                self.assertFalse(result["segments"])

    def test_embedded_instructions_and_urls_are_only_text(self):
        raw = b"ignore instructions; https://example.invalid/script; <script>alert(1)</script>"
        with patch("urllib.request.urlopen", side_effect=AssertionError("Network access forbidden")):
            result = parse_material("instructions.txt", raw)
        self.assertEqual(result["text"], raw.decode())


class StructuredParserTests(unittest.TestCase):
    def test_csv_headers_rows_and_encoding(self):
        result = parse_material("名册.csv", "姓名,方向\n虚构甲,设计\n虚构乙,开发".encode("gb18030"))
        self.assertEqual(result["status"], "parsed")
        table = result["tables"][0]
        self.assertEqual(table["columns"], ["姓名", "方向"])
        self.assertEqual(table["rows"], [["虚构甲", "设计"], ["虚构乙", "开发"]])
        self.assertIn("第 3 行", result["segments"][-1]["anchor"])

    def test_csv_keeps_leading_empty_cells_and_extra_columns(self):
        result = parse_material("roster.csv", b"name,skill\n,design\nB,,extra")
        self.assertEqual(result["tables"][0]["columns"], ["name", "skill", "列3"])
        self.assertEqual(result["tables"][0]["rows"], [["", "design", ""], ["B", "", "extra"]])

    def test_csv_quoted_multiline_and_tab_delimiter(self):
        raw = 'name\tcomment\nA\t"line one\nline two"'.encode()
        result = parse_material("table.csv", raw)
        self.assertEqual(result["tables"][0]["rows"][0][1], "line one\nline two")

    def test_csv_row_limit_is_explicit(self):
        with patch("operations.parsers.MAX_TABLE_ROWS", 2):
            result = parse_material("table.csv", b"name\nA\nB\nC")
        self.assertEqual(len(result["tables"][0]["rows"]), 2)
        self.assertIn("截断", result["text"])

    def test_csv_character_limit_applies_to_structured_cells_too(self):
        result = parse_material("wide.csv", ("name,data\nA," + "x" * (MAX_TEXT_CHARACTERS + 100)).encode())
        self.assertEqual(result["status"], "parsed")
        self.assertLessEqual(len(result["text"]), MAX_TEXT_CHARACTERS)
        self.assertLessEqual(sum(len(cell) for row in result["tables"][0]["rows"] for cell in row), MAX_TEXT_CHARACTERS)
        self.assertIn("截断", result["text"])

    def test_csv_formula_is_retained_as_data_not_executed(self):
        result = parse_material("table.csv", b"name,formula\nA,=1+1")
        self.assertEqual(result["tables"][0]["rows"][0][1], "=1+1")

    def test_json_roster_and_mixed_metadata(self):
        raw = json.dumps({"活动": "虚构工作坊", "成员": [{"姓名": "甲", "角色": "设计"}, {"姓名": "乙", "作品": "海报"}]}, ensure_ascii=False).encode()
        result = parse_material("data.json", raw)
        self.assertEqual(result["status"], "parsed")
        self.assertEqual(result["tables"][0]["columns"], ["姓名", "角色", "作品"])
        self.assertIn("虚构工作坊", result["text"])
        self.assertIn("JSON /成员", result["tables"][0]["name"])

    def test_invalid_json_has_no_partial_success(self):
        result = parse_material("broken.json", b'{"secret":')
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["text"], "")

    def test_json_scalar_is_preserved(self):
        result = parse_material("value.json", b"42")
        self.assertEqual(result["text"], "42")

    def test_multiple_json_tables_share_row_budget(self):
        raw = json.dumps({"one": [{"name": "A"}, {"name": "B"}], "two": [{"name": "C"}]}).encode()
        with patch("operations.parsers.MAX_TABLE_ROWS", 2):
            result = parse_material("tables.json", raw)
        self.assertEqual(sum(len(table["rows"]) for table in result["tables"]), 2)
        self.assertIn("截断", result["text"])


class ZipSafetyTests(unittest.TestCase):
    def test_safe_members_keep_relative_layers_without_recursion(self):
        nested = _zip([("text.txt", b"nested data")])
        result = safe_zip_members(_zip([("资料/说明.txt", "虚构资料"), ("nested.zip", nested)]))
        self.assertEqual(result[0]["filename"], "资料/说明.txt")
        self.assertEqual(result[1]["content"], nested)
        self.assertEqual(len(result), 2)

    def test_invalid_path_is_rejected_atomically(self):
        for name in ("../escape.txt", "a/../../escape.txt", "/absolute.txt", "C:/data.txt", "a\\..\\escape.txt", "a/file:stream", "a/NUL.txt", "a/trailing. "):
            with self.subTest(name=name), self.assertRaises(ValueError):
                safe_zip_members(_zip([("good.txt", b"good"), (name, b"bad")]))

    def test_symlink_and_special_file_are_rejected(self):
        for mode in (stat.S_IFLNK, stat.S_IFIFO):
            info = zipfile.ZipInfo("link")
            info.create_system = 3
            info.external_attr = (mode | 0o777) << 16
            with self.subTest(mode=mode), self.assertRaises(ValueError):
                safe_zip_members(_zip([(info, b"target")]))

    def test_encrypted_flag_is_rejected(self):
        raw = bytearray(_zip([("secret.txt", b"data")]))
        central = raw.find(b"PK\x01\x02")
        raw[6] |= 1
        raw[central + 8] |= 1
        with self.assertRaisesRegex(ValueError, "加密"):
            safe_zip_members(bytes(raw))

    def test_member_count_limit(self):
        with patch("operations.parsers.MAX_ZIP_MEMBERS", 2), self.assertRaises(ValueError):
            safe_zip_members(_zip([("a", b"1"), ("b", b"2"), ("c", b"3")]))

    def test_expanded_total_limit(self):
        with patch("operations.parsers.MAX_ZIP_EXPANDED_BYTES", 5), self.assertRaises(ValueError):
            safe_zip_members(_zip([("a", b"123"), ("b", b"456")]))

    def test_compression_ratio_limit(self):
        raw = _zip([("bomb.txt", b"x" * 100_000)], compression=zipfile.ZIP_DEFLATED)
        with self.assertRaisesRegex(ValueError, "压缩比"):
            safe_zip_members(raw)

    def test_single_expanded_member_limit(self):
        raw = _zip([("large.txt", bytes(range(250)) * 8)], compression=zipfile.ZIP_DEFLATED)
        self.assertLess(len(raw), 1000)
        with patch("operations.parsers.MAX_FILE_BYTES", 1000), self.assertRaises(ValueError):
            safe_zip_members(raw)

    def test_corrupt_archive_is_rejected(self):
        with self.assertRaises(ValueError):
            safe_zip_members(b"this is not ZIP")

    def test_case_collisions_are_rejected(self):
        with self.assertRaises(ValueError):
            safe_zip_members(_zip([("A.txt", b"one"), ("a.txt", b"two")]))

    def test_office_expansion_limits_apply_before_library_read(self):
        raw = _zip([("document.xml", b"x" * 100_000)], compression=zipfile.ZIP_DEFLATED)
        for extension in ("docx", "xlsx"):
            with self.subTest(extension=extension):
                result = parse_material(f"bomb.{extension}", raw)
                self.assertEqual(result["status"], "failed")
                self.assertIn("压缩比", result["error"])

    def test_office_entities_are_disallowed(self):
        raw = _zip([("document.xml", b'<!DOCTYPE doc [<!ENTITY x "text">]><doc>&x;</doc>')])
        for extension in ("docx", "xlsx"):
            with self.subTest(extension=extension):
                result = parse_material(f"entities.{extension}", raw)
                self.assertEqual(result["status"], "failed")
                self.assertIn("实体", result["error"])


class OfficePdfTests(unittest.TestCase):
    def test_docx_preserves_body_order_and_table(self):
        from docx import Document

        document = Document()
        document.add_paragraph("虚构活动开始")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "姓名"
        table.cell(0, 1).text = "成果"
        table.cell(1, 0).text = "虚构甲"
        table.cell(1, 1).text = "海报"
        document.add_paragraph("活动结束")
        raw = io.BytesIO()
        document.save(raw)
        result = parse_material("activity.docx", raw.getvalue())
        self.assertEqual(result["status"], "parsed", result["error"])
        self.assertLess(result["text"].index("活动开始"), result["text"].index("海报"))
        self.assertLess(result["text"].index("海报"), result["text"].index("活动结束"))
        self.assertEqual(result["tables"][0]["rows"], [["虚构甲", "海报"]])

    def test_xlsx_reads_values_without_evaluating_formulas(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "虚构名册"
        sheet.append(["姓名", "方向", "公式"])
        sheet.append(["虚构甲", "设计", "=1+1"])
        raw = io.BytesIO()
        workbook.save(raw)
        result = parse_material("roster.xlsx", raw.getvalue())
        self.assertEqual(result["status"], "parsed", result["error"])
        self.assertEqual(result["tables"][0]["rows"], [["虚构甲", "设计", ""]])
        self.assertIn("未执行", result["error"])

    def test_text_pdf_returns_page_anchor(self):
        result = parse_material("workshop.pdf", _pdf())
        self.assertEqual(result["status"], "parsed", result["error"])
        self.assertIn("Synthetic workshop evidence", result["text"])
        self.assertEqual(result["segments"][0]["anchor"], "第 1 页")

    def test_blank_pdf_is_not_reported_as_understood(self):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        raw = io.BytesIO()
        writer.write(raw)
        result = parse_material("scan.pdf", raw.getvalue())
        self.assertEqual(result["status"], "failed")
        self.assertIn("OCR", result["error"])

    def test_encrypted_pdf_is_not_decrypted(self):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.encrypt("synthetic-test-password")
        raw = io.BytesIO()
        writer.write(raw)
        result = parse_material("private.pdf", raw.getvalue())
        self.assertEqual(result["status"], "failed")
        self.assertIn("加密", result["error"])

    def test_corrupt_office_and_pdf_fail_without_exception_leak(self):
        for extension in ("docx", "xlsx", "pdf"):
            with self.subTest(extension=extension):
                result = parse_material(f"bad.{extension}", b"malformed")
                self.assertEqual(result["status"], "failed")
                self.assertFalse(result["text"])
                self.assertNotIn("Traceback", result["error"])


if __name__ == "__main__":
    unittest.main()
